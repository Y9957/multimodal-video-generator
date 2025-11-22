"""
ppt_parser.py
- PPTX에서 텍스트 / 이미지 / 표 추출
- 슬라이드 PNG 스냅샷 생성
- SlideData 및 State
"""

import os
import re
import subprocess
from pathlib import Path
from typing import List, Dict, Optional, TypedDict
from dataclasses import dataclass

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# ------------------------------------------------------------
# SlideData / State 정의 
# ------------------------------------------------------------

@dataclass
class SlideData:
    page: int                          # 페이지 번호
    slide_image: str                   # 페이지 스냅샷 이미지
    texts: List[str]                   # 페이지 텍스트
    images: List[str]                  # 페이지 이미지
    tables: List[List[str]]            # 페이지 표
    summary: Optional[str] = None      # 페이지 요약
    script: Optional[str] = None       # 강의 스크립트
    script_file: Optional[str] = None  # 스크립트 파일 경로
    audio: Optional[str] = None        # 음성 파일 경로
    video: Optional[str] = None        # 비디오 파일 경로


class State(TypedDict, total=False):
    # 입력 / 기본 정보
    pptx_path: str                     # PPTX 경로
    prompt: Dict[str, str]             # 사용자 음성/스타일 프롬프트
    work_dir: str                      # 작업 폴더
    media_dir: str                     # 미디어 출력 폴더

    # 추출 산출물
    slides: List[SlideData]            # 페이지 파싱 결과

    # 생성 산출물
    full_script_path: str              # 전체 스크립트 파일 경로

    # 미디어 산출물
    full_video_path: str               # 최종 결합 영상 경로


# ------------------------------------------------------------
# 유틸 함수
# ------------------------------------------------------------

def clean_text(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


# ------------------------------------------------------------
# PPT → PNG 스냅샷 
# ------------------------------------------------------------

def export_slide_as_png(state: State, idx: int, dpi: int = 220) -> str:
    """
    PPTX → PNG 변환 (LibreOffice + pdftoppm)
    원본 코드 로직 유지
    """
    work_dir = Path(state["work_dir"]).expanduser().resolve()
    work_dir.mkdir(parents=True, exist_ok=True)

    pptx = Path(state["pptx_path"]).expanduser().resolve()
    if not pptx.exists():
        raise FileNotFoundError(f"PPTX 없음: {pptx}")

    page_no = idx + 1
    out_prefix = work_dir / "slide_img"

    env = os.environ.copy()
    env.update({"LANG": "ko_KR.UTF-8", "LC_ALL": "ko_KR.UTF-8"})

    # 1) libreoffice → png
    before_png = set(work_dir.glob("*.png"))
    subprocess.run(
        [
            "soffice", "--headless",
            "-env:UserInstallation=file:///tmp/lo_profile",
            "--convert-to", "png:impress_png_Export",
            "--outdir", str(work_dir),
            str(pptx),
        ],
        capture_output=True, text=True, env=env
    )

    created_png = [p for p in work_dir.glob("*.png") if p not in before_png]

    # 정확한 페이지 찾기
    match_png = [p for p in created_png if p.stem.endswith(f"-{page_no}")]
    if match_png:
        png_path = max(match_png, key=lambda p: p.stat().st_mtime)
        return str(png_path)

    # 2) fallback: pdf → png
    pdf_path = work_dir / f"{pptx.stem}.pdf"
    subprocess.run(
        [
            "soffice", "--headless",
            "-env:UserInstallation=file:///tmp/lo_profile",
            "--convert-to", "pdf:impress_pdf_Export",
            "--outdir", str(work_dir),
            str(pptx),
        ],
        capture_output=True, text=True, env=env
    )

    # pdf → png
    subprocess.run(
        [
            "pdftoppm",
            "-f", str(page_no),
            "-l", str(page_no),
            "-png", "-r", str(dpi),
            str(pdf_path),
            str(out_prefix),
        ],
        capture_output=True, text=True, env=env
    )

    fallback_path = Path(f"{out_prefix}-{page_no}.png")
    return str(fallback_path)


# ------------------------------------------------------------
# node_parse_ppt (핵심 함수) 
# ------------------------------------------------------------

def node_parse_ppt(state: State) -> State:
    """
    PPTX의 각 페이지에서 텍스트/이미지/표 추출 후 SlideData로 저장
    PNG 스냅샷도 생성
    """

    prs = Presentation(state["pptx_path"])
    slides: List[SlideData] = []

    for i, slide in enumerate(prs.slides):
        texts, images, tables = [], [], []

        for shape in slide.shapes:

            # 텍스트 추출
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    t = clean_text(paragraph.text)
                    if t:
                        texts.append(t)

            # 표 추출
            elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                table_data = []
                for row in shape.table.rows:
                    table_data.append([clean_text(cell.text) for cell in row.cells])
                tables.append(table_data)

            # 이미지 추출
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                ext = img.ext or "png"
                filename = f"{state['media_dir']}/slide{i}_img{len(images)+1}.{ext}"
                with open(filename, "wb") as f:
                    f.write(img.blob)
                images.append(filename)

        # 슬라이드 이미지 생성
        slide_image = export_slide_as_png(state, i)

        slide_data = SlideData(
            page=i,
            slide_image=slide_image,
            texts=texts,
            images=images,
            tables=tables
        )
        slides.append(slide_data)

        print(f"[INFO] Slide {i}: 텍스트 {len(texts)}, 이미지 {len(images)}, 표 {len(tables)}")

    state["slides"] = slides
    print(f"[INFO] 총 {len(slides)}개 슬라이드 파싱 완료.")
    return state
